import pathlib, sys, json
ROOT = pathlib.Path(__file__).resolve().parent.parent

import pymupdf
from docx import Document
import openpyxl
from pptx import Presentation
from pptx.util import Inches

def make(dest):
    dest = pathlib.Path(dest)
    dest.mkdir(parents=True, exist_ok=True)
    text = '验收编号 4729。计划：周五交付。预算：128 元。'
    (dest / 'sample.txt').write_text(text, encoding='utf-8')
    (dest / 'sample-gbk.txt').write_bytes(text.encode('gb18030'))
    doc = pymupdf.open()
    page = doc.new_page(width=700, height=350)
    page.insert_text((35, 100), text, fontname='china-s', fontsize=20)
    doc.save(dest / 'sample.pdf')
    pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))
    scan = pymupdf.open()
    scan.new_page(width=700, height=350).insert_image(pymupdf.Rect(0, 0, 700, 350), stream=pix.tobytes('png'))
    scan.save(dest / 'scan.pdf')
    scan.close()
    doc.close()
    word = Document()
    word.add_paragraph(text)
    word.save(dest / 'sample.docx')
    book = openpyxl.Workbook()
    book.active.append(['编号', '预算'])
    book.active.append([4729, 128])
    book.save(dest / 'sample.xlsx')
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.text = text
    slides.save(dest / 'sample.pptx')
    return {'created': ['sample.txt', 'sample-gbk.txt', 'sample.pdf', 'scan.pdf', 'sample.docx', 'sample.xlsx', 'sample.pptx']}

if __name__ == '__main__':
    print(json.dumps(make(sys.argv[1]), ensure_ascii=False))
