"""Bounded local attachment extraction; no macros, external links or shell execution."""
import sys, pathlib, json, zipfile, io

ROOT = pathlib.Path(__file__).resolve().parent.parent

def read_document(file, options):
    path = pathlib.Path(file)
    suffix = path.suffix.lower()
    limit = min(40000, int(options.get('maxChars', 20000)))
    pages_limit = min(100, int(options.get('maxPages', 50)))
    image_limit = min(8, int(options.get('scanPages', 6)))
    result = {'text': '', 'images': [], 'truncated': False}
    if suffix in ('.docx', '.xlsx', '.pptx'):
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            if len(infos) > 10000 or sum(i.file_size for i in infos) > 128 * 1024 * 1024:
                raise ValueError('压缩后的 Office 文档展开体积过大')
    if suffix == '.pdf':
        import pymupdf
        with pymupdf.open(path) as doc:
            if doc.needs_pass:
                raise ValueError('PDF 已加密，请先发送解密后的文件')
            result['pages'] = len(doc)
            result['truncated'] = len(doc) > pages_limit
            out = pathlib.Path(options['outputDir'])
            out.mkdir(parents=True, exist_ok=True)
            chunks = []
            scanned = 0
            for i, page in enumerate(doc):
                if i >= pages_limit or sum(map(len, chunks)) >= limit:
                    result['truncated'] = True
                    break
                text = page.get_text(sort=True).strip()
                chunks.append(f'第 {i+1} 页：\n{text}')
                # Scanned pages go to the existing visual model, with their page number.
                if len(text) < 30:
                    scanned += 1
                    if len(result['images']) < image_limit:
                        scale = min(1.8, 1400 / max(page.rect.width, page.rect.height))
                        dest = out / f'page-{i+1}.png'
                        page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False).save(dest)
                        result['images'].append({'file': str(dest), 'page': i+1})
            result['text'] = '\n\n'.join(chunks)
            if scanned > len(result['images']):
                result['truncated'] = True
    elif suffix == '.docx':
        from docx import Document
        doc = Document(path)
        chunks = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            chunks.extend('\t'.join(cell.text for cell in row.cells) for row in table.rows)
        result['text'] = '\n'.join(chunks)
    elif suffix == '.xlsx':
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        chunks = []
        try:
            for sheet in wb.worksheets[:10]:
                chunks.append('工作表：' + sheet.title)
                for n, row in enumerate(sheet.iter_rows(values_only=True)):
                    if n >= 500 or sum(map(len, chunks)) >= limit:
                        result['truncated'] = True
                        break
                    chunks.append('\t'.join('' if x is None else str(x) for x in row[:50]))
        finally:
            wb.close()
        result['text'] = '\n'.join(chunks)
    elif suffix == '.pptx':
        from pptx import Presentation
        pres = Presentation(path)
        chunks = []
        for i, slide in enumerate(pres.slides):
            if i >= pages_limit:
                result['truncated'] = True
                break
            chunks.append('幻灯片 ' + str(i+1))
            for shape in slide.shapes:
                if shape.has_text_frame:
                    chunks.append(shape.text)
                if shape.has_table:
                    chunks.extend('\t'.join(c.text for c in row.cells) for row in shape.table.rows)
        result['text'] = '\n'.join(chunks)
    else:
        raise ValueError('暂不支持这个文档格式')
    result['truncated'] = result['truncated'] or len(result['text']) > limit
    result['text'] = result['text'][:limit]
    return result

if __name__ == '__main__':
    try:
        print(json.dumps(read_document(sys.argv[1], json.loads(sys.argv[2])), ensure_ascii=False))
    except Exception as e:
        # No document content or local path in runtime failure output.
        print(json.dumps({'error': str(e) if isinstance(e, ValueError) else '文件损坏或文档读取失败'}, ensure_ascii=False))
        sys.exit(1)
